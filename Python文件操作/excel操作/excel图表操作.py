# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.chart import (
    Reference,
    Series,
    PieChart,
    BarChart,
    BubbleChart
)

wb = openpyxl.Workbook()

# Pie chart
ws = wb.active
ws.title = 'pie_chart'
data = [
    ['pie', 'sold'],
    ['apple', 50],
    ['cherry', 30],
    ['pumpkin', 10],
    ['chocolate', 40]
]

for row in data:
    ws.append(row)

pie = PieChart()
labels = Reference(ws, min_col=1, min_row=2, max_row=5)
data = Reference(ws, min_col=2, min_row=2, max_row=5)

pie.add_data(data)
pie.set_categories(labels)
pie.title = 'Pies sold by category'
ws.add_chart(pie, 'A15')

# Bar chart
ws = wb.create_sheet('bar_chart')

rows = [
    ('Number', 'Batch 1', 'Batch 2'),
    (2, 10, 30),
    (3, 40, 50),
    (4, 50, 70),
    (5, 20, 10),
    (6, 10, 40),
    (7, 50, 30)
]

for row in rows:
    ws.append(row)

chart1 = BarChart()
# bar的两种图标：一个是水平，一个垂直
chart1.type = 'bar'     # 水平
# chart1.type = 'col'   # 垂直
chart1.style = 15
chart1.title = 'Bar Chart'
chart1.y_axis.title = 'Sample length(mm)'
chart1.x_axis.title = 'Test number'

# 使用单列值时，就不用说明最大列了
cates = Reference(ws, min_col=1, min_row=2, max_row=7)
# 多行多列的值，都给指出范围
data = Reference(ws, min_col=2, max_col=3, min_row=2, max_row=7)
chart1.add_data(data)
chart1.set_categories(cates)
ws.add_chart(chart1, 'A10')

wb.save('excel_chart.xlsx')
